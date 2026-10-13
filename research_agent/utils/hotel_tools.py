"""hotel-checker 的 4 个工具（plan.md 7.3），分两层写。

上半部分是逻辑层：普通函数，返回 Python 数据，出错时抛 ValueError。
下半部分是 @tool 包装层：把结果转成 AI 读的文字，把 ValueError 转成 is_error。
数据存在内存里，每家酒店一份，不写中间文件；最后只输出 xlsx 和 summary.csv。
逻辑从原 skill 的 add_keyword.py、log_round.py、build_report.py 改过来。
"""

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from claude_agent_sdk import create_sdk_mcp_server, tool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

# 第 0 轮的核心词，按正则表达式匹配。用 wall 不用 thin wall：
# "walls are thin"、"non soundproof walls" 也能搜到；搜宽了只是多读几条
CORE_KEYWORD = "nois|wall|hear|siren"
# 第 1～7 轮各加一个词，必须跑满 7 轮
MAX_ROUND = 7
# next_batch 每次最多给几条
BATCH = 50
# 判定规则（用户定的）
OCC_LIMIT, UNAV_LIMIT = 0.2, 0.3
OK_VERDICT = "OK to book"
NOT_OK_VERDICT = "Not suitable for noise-sensitive guests"
CATEGORIES = ("unavoidable", "avoidable", "unclear")
REPORT = "hotel_noise_analysis_log.xlsx"
SUMMARY_HEADER = ["city", "score", "hotel", "verdict", "occurrence_rate",
                  "unavoidable_rate", "rounds", "report"]
LOG_HEADER = ["round", "keyword", "reasoning", "new_hits_current_round",
              "new_positive_noise_reviews", "occurrence_rate",
              "unavoidable_rate", "hit_rows", "noise_rows"]


@dataclass
class Review:
    """一条写了差评的评论，和 agent 对它的判断。"""
    row: int                # 原数据行号
    text: str               # 差评原文（已整理空格）
    round: int | None = None  # 第几轮被关键词命中；None = 还没命中过
    is_noise: str | None = None  # "yes" / "no"；None = 还没判断
    category: str = ""
    quote: list[str] = field(default_factory=list)
    reason: str = ""


@dataclass
class Hotel:
    """一家酒店在内存里的全部数据。"""
    name: str
    work_dir: Path          # runs/<城市>/<评分>/<酒店文件夹名>
    total: int              # 全部评论条数，含没写差评的
    reviews: list[Review]   # 写了差评的评论，按行号排
    # 每轮一项 (keyword, reasoning)；列表下标就是轮次，第 0 轮是核心词
    rounds: list[tuple[str, str]]


# 全部酒店：酒店名 -> Hotel。由 load_hotels 填好，4 个工具读写它
HOTELS: dict[str, Hotel] = {}


def mark_round(hotel, pattern, round_no):
    """还没命中过、而且含 pattern 的差评，标成第 round_no 轮。返回标了几条。"""
    count = 0
    for review in hotel.reviews:
        if review.round is None and pattern.search(review.text):
            review.round = round_no
            count += 1
    return count


def load_hotels(hotels):
    """把 select_hotels 给的酒店存进内存（先清空旧的），并用核心词标出第 0 轮。"""
    HOTELS.clear()
    for h in hotels:
        hotel = Hotel(
            name=h["hotel"],
            work_dir=Path(h["work_dir"]),
            total=h["reviews"],
            reviews=[Review(c["row"], c["text"]) for c in h["complaints"]],
            rounds=[(CORE_KEYWORD, "initial core keywords")],
        )
        # re.I = 不分大小写
        mark_round(hotel, re.compile(CORE_KEYWORD, re.I), 0)
        HOTELS[hotel.name] = hotel


def get_hotel(name):
    """按酒店名找数据；名字必须和选中的酒店完全一致，抄错了就当场报错。"""
    if name not in HOTELS:
        raise ValueError(f"Unknown hotel: {name!r}. "
                         "Use the hotel name exactly as you were given it.")
    return HOTELS[name]


def current_round(hotel):
    """现在是第几轮：第 0 轮之后每加一个词加 1。"""
    return len(hotel.rounds) - 1


def unjudged(hotel):
    """本轮命中、还没判断的差评。"""
    r = current_round(hotel)
    return [v for v in hotel.reviews if v.round == r and v.is_noise is None]


def next_batch(name):
    """本轮还没判断的下一批（最多 BATCH 条）；本轮读完时只给条数，不给两个比例。"""
    hotel = get_hotel(name)
    r = current_round(hotel)
    left = unjudged(hotel)
    if left:
        return {"round": r, "done": False, "left": len(left),
                "reviews": [{"row": v.row, "text": v.text}
                            for v in left[:BATCH]]}
    this_round = [v for v in hotel.reviews if v.round == r]
    return {"round": r, "done": True, "read": len(this_round),
            "noise": sum(v.is_noise == "yes" for v in this_round)}


def normalize(text):
    """转小写、去标点、多个空格合成一个（quote 检查用）。"""
    return " ".join(re.sub(r"[^\w\s]", " ", text.lower()).split())


def quote_found(quote, text):
    """quote 整理后不能是空的，并且能在整理后的原文里找到。"""
    piece = normalize(quote)
    return bool(piece) and piece in normalize(text)


def category_errors(is_noise, category):
    """yes 时 category 只能三选一；no 时必须留空。"""
    if is_noise == "yes" and category not in CATEGORIES:
        return [f"category must be one of {', '.join(CATEGORIES)}, "
                f"got {category!r}"]
    if is_noise == "no" and category:
        return ["category must be empty when is_noise is 'no'"]
    return []


def quote_errors(is_noise, quote, text):
    """quote 是一个列表；yes 时至少 1 句、每句都出自原文；no 时必须为空。"""
    if not isinstance(quote, list) or \
            not all(isinstance(q, str) for q in quote):
        return ["quote must be a list of sentences"]
    if is_noise == "yes" and not quote:
        return ["quote needs at least one sentence when is_noise is 'yes'"]
    if is_noise == "no" and quote:
        return ["quote must be empty when is_noise is 'no'"]
    return [f"quote not found in the review: {q!r}"
            for q in quote if not quote_found(q, text)]


def item_errors(item, review):
    """检查一条判断，返回问题清单（空 = 没问题），每条都带行号。"""
    is_noise = item.get("is_noise")
    errors = []
    if is_noise not in ("yes", "no"):
        errors.append(f"is_noise must be 'yes' or 'no', got {is_noise!r}")
    errors += category_errors(is_noise, item.get("category", ""))
    errors += quote_errors(is_noise, item.get("quote"), review.text)
    reason = item.get("reason")
    if not isinstance(reason, str) or not reason.strip():
        errors.append("reason is empty")
    return [f"row {review.row}: {e}" for e in errors]


def batch_errors(hotel, items):
    """检查一整批判断：行号属于本轮、没判断过、没重复，再逐条检查内容。"""
    r = current_round(hotel)
    this_round = {v.row: v for v in hotel.reviews if v.round == r}
    errors, seen = [], set()
    for item in items:
        row = item.get("row")
        review = this_round.get(row)
        if row in seen:
            errors.append(f"row {row}: sent twice in this batch")
        elif review is None:
            errors.append(f"row {row}: not a review of round {r}")
        elif review.is_noise is not None:
            errors.append(f"row {row}: already judged")
        else:
            errors += item_errors(item, review)
        seen.add(row)
    return errors


def add_judgments(name, items):
    """存一批判断；有一条不对整批退回，一条都不存。存完返回下一批。"""
    hotel = get_hotel(name)
    if not items:
        raise ValueError("items is empty: send at least one judgment")
    errors = batch_errors(hotel, items)
    if errors:
        raise ValueError("Nothing was saved. Fix these and send the whole "
                         "batch again:\n" + "\n".join(errors))
    by_row = {v.row: v for v in hotel.reviews}
    for item in items:
        review = by_row[item["row"]]
        review.is_noise = item["is_noise"]
        review.category = item.get("category", "") if item["is_noise"] == "yes" else ""
        review.quote = list(item["quote"])
        review.reason = item["reason"].strip()
    return next_batch(name)


def add_keyword(name, word, reasoning):
    """上一轮读完、没满 7 轮时加一个词：含新词、还没命中过的差评标为新一轮。"""
    hotel = get_hotel(name)
    r = current_round(hotel)
    if unjudged(hotel):
        raise ValueError(f"Round {r} is not finished: judge all its "
                         "reviews first (call next_batch)")
    if r >= MAX_ROUND:
        raise ValueError(f"All {MAX_ROUND} keyword rounds are done: "
                         "call build_report")
    word, reasoning = word.strip(), reasoning.strip()
    if not word:
        raise ValueError("word is empty")
    if not reasoning:
        raise ValueError("reasoning is empty")
    hotel.rounds.append((word, reasoning))
    # re.escape：词里的特殊符号按普通文字处理；re.I = 不分大小写
    new_hits = mark_round(hotel, re.compile(re.escape(word), re.I), r + 1)
    return {"new_hits": new_hits, **next_batch(name)}


def rates(reviews, total):
    """一组评论 -> (occurrence_rate, unavoidable_rate)。

    occurrence_rate = 噪音差评条数 / 评论总数；
    unavoidable_rate = unavoidable 条数 / 噪音差评条数，还没有噪音差评时算 0。
    """
    noise = [v for v in reviews if v.is_noise == "yes"]
    unavoidable = sum(v.category == "unavoidable" for v in noise)
    occ = len(noise) / total
    unav = unavoidable / len(noise) if noise else 0.0
    return occ, unav


def verdict_of(occ, unav):
    """两个比例都低于门槛才能订。"""
    if occ < OCC_LIMIT and unav < UNAV_LIMIT:
        return OK_VERDICT
    return NOT_OK_VERDICT


def log_rows(hotel):
    """log 表：每轮一行，比例算到这一轮为止（现算，不存快照）。"""
    rows = []
    for r, (keyword, reasoning) in enumerate(hotel.rounds):
        hits = [v for v in hotel.reviews if v.round == r]
        noise = [v for v in hits if v.is_noise == "yes"]
        so_far = [v for v in hotel.reviews
                  if v.round is not None and v.round <= r]
        occ, unav = rates(so_far, hotel.total)
        rows.append([r, keyword, reasoning, len(hits), len(noise), occ, unav,
                     ", ".join(str(v.row) for v in hits),
                     ", ".join(str(v.row) for v in noise)])
    return rows


def sheet(ws, header, data, pct=(), wide=()):
    """写一个工作表：表头加粗，比例显示成百分比，长文字列加宽并换行。"""
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for d in data:
        ws.append(d)
    for i, h in enumerate(header, 1):
        # 列字母，如 A、B
        col = ws.cell(1, i).column_letter
        ws.column_dimensions[col].width = 60 if h in wide else 18
        for c in ws[col][1:]:
            if h in pct:
                c.number_format = "0.0%"
            if h in wide:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    # 往下滚动时表头不动
    ws.freeze_panes = "A2"


def write_xlsx(hotel, result, inference, path):
    """生成 4 个工作表的 Excel：summary、log、judgments、all negative reviews。"""
    judged = [v for v in hotel.reviews if v.is_noise is not None]
    noise = [v for v in judged if v.is_noise == "yes"]
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    sheet(ws, ["item", "value"], [
        ["hotel", hotel.name],
        ["verdict", result["verdict"]],
        ["total_reviews", hotel.total],
        ["noise_reviews", len(noise)],
        ["unavoidable_reviews",
         sum(v.category == "unavoidable" for v in noise)],
        ["occurrence_rate", result["occurrence_rate"]],
        ["unavoidable_rate", result["unavoidable_rate"]],
        ["keywords_used", "; ".join(k for k, _ in hotel.rounds)],
        ["inference", inference],
    ], wide=("value",))
    # 两个指标显示成百分比
    for row in ws.iter_rows(min_row=2):
        if row[0].value in ("occurrence_rate", "unavoidable_rate"):
            row[1].number_format = "0.0%"

    sheet(wb.create_sheet("log"), LOG_HEADER, log_rows(hotel),
          pct=("occurrence_rate", "unavoidable_rate"),
          wide=("reasoning", "hit_rows", "noise_rows"))
    # 每条已读差评的判断，按原数据行号排；多句 quote 用 | 连起来显示
    sheet(wb.create_sheet("judgments"),
          ["row", "round", "text", "is_noise", "category", "quote", "reason"],
          [[v.row, v.round, v.text, v.is_noise, v.category,
            " | ".join(v.quote), v.reason] for v in judged],
          wide=("text", "quote", "reason"))
    # 所有写了差评的评论；query_hit = 命中核心词（第 0 轮）
    sheet(wb.create_sheet("all negative reviews"),
          ["row", "query_hit", "text"],
          [[v.row, "Y" if v.round == 0 else "", v.text]
           for v in hotel.reviews],
          wide=("text",))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_summary(hotel, result, report):
    """往城市总表 runs/<城市>/summary.csv 写这家酒店一行；同一家的旧行先删掉。

    城市、评分直接用文件夹名（runs/<城市>/<评分>/<酒店>），和文件夹一致；
    report 写相对 runs/<城市>/ 的路径，不带电脑用户名。
    """
    city_dir = hotel.work_dir.parent.parent
    path = city_dir / "summary.csv"
    rows = []
    if path.exists():
        with open(path, newline="", encoding="utf-8") as f:
            rows = [r for r in csv.reader(f)
                    if r and r != SUMMARY_HEADER and r[2] != hotel.name]
    rows.append([city_dir.name, hotel.work_dir.parent.name, hotel.name,
                 result["verdict"], round(result["occurrence_rate"], 4),
                 round(result["unavoidable_rate"], 4), result["rounds"],
                 report.relative_to(city_dir).as_posix()])
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(SUMMARY_HEADER)
        writer.writerows(rows)


def build_report(name, inference):
    """跑满 7 轮后：现算指标和结论，生成 xlsx，写 summary.csv，返回 6 项结果。"""
    hotel = get_hotel(name)
    r = current_round(hotel)
    if r < MAX_ROUND:
        raise ValueError(f"Only {r} of {MAX_ROUND} keyword rounds are done: "
                         f"all {MAX_ROUND} are required (call add_keyword)")
    if unjudged(hotel):
        raise ValueError(f"Round {r} is not finished: judge all its "
                         "reviews first (call next_batch)")
    inference = inference.strip()
    if not inference:
        raise ValueError("inference is empty")
    occ, unav = rates(hotel.reviews, hotel.total)
    path = hotel.work_dir / REPORT
    result = {"hotel": hotel.name, "verdict": verdict_of(occ, unav),
              "occurrence_rate": occ, "unavoidable_rate": unav,
              "rounds": len(hotel.rounds), "report": str(path)}
    write_xlsx(hotel, result, inference, path)
    write_summary(hotel, result, path)
    return result


# ---------------- @tool 包装层 ----------------

# mcp_servers 里的键名；AI 看到的工具全名是 mcp__hotel__<工具名>
SERVER_NAME = "hotel"
HOTEL_PARAM = {"type": "string",
               "description": "The hotel name, exactly as you were given it"}


def batch_text(result):
    """next_batch 的结果 -> AI 读的文字：这一批差评，或本轮读完 + 下一步做什么。"""
    r = result["round"]
    if not result["done"]:
        lines = [f"Round {r}: {result['left']} review(s) left to judge. "
                 f"This batch has {len(result['reviews'])}:"]
        lines += [f"row {v['row']}: {v['text']}" for v in result["reviews"]]
        lines.append("Judge every review above, then send them all "
                     "with add_judgments.")
        return "\n".join(lines)
    done = (f"Round {r} finished: read {result['read']} review(s), "
            f"{result['noise']} judged as noise.")
    if r < MAX_ROUND:
        return f"{done} Next: call add_keyword (round {r + 1} of {MAX_ROUND})."
    return (f"{done} All {MAX_ROUND} keyword rounds are finished. "
            "Next: write the inference and call build_report.")


def report_text(result):
    """build_report 的结果 -> 6 行，hotel-checker 原样回复给 lead。"""
    return "\n".join([
        f"hotel: {result['hotel']}",
        f"verdict: {result['verdict']}",
        f"occurrence_rate: {result['occurrence_rate']:.1%}",
        f"unavoidable_rate: {result['unavoidable_rate']:.1%}",
        f"rounds: {result['rounds']}",
        f"report: {result['report']}",
    ])


def run(func, to_text, *args):
    """调用逻辑层：成功就转成文字；ValueError 转成 is_error，AI 看到原因后改了重交。"""
    try:
        text = to_text(func(*args))
    except ValueError as error:
        return {"content": [{"type": "text", "text": str(error)}],
                "is_error": True}
    return {"content": [{"type": "text", "text": text}]}


@tool("next_batch",
      "Get the next batch (up to 50) of unjudged reviews in the current "
      "round, or the news that the round is finished.",
      {"type": "object", "properties": {"hotel": HOTEL_PARAM},
       "required": ["hotel"]})
async def next_batch_tool(args):
    """工具 next_batch。"""
    return run(next_batch, batch_text, args["hotel"])


@tool("add_judgments",
      "Save your judgments for reviews of the current round. If any item is "
      "wrong, nothing is saved and the error says what to fix. Returns the "
      "next batch.",
      {"type": "object",
       "properties": {
           "hotel": HOTEL_PARAM,
           "items": {
               "type": "array",
               "minItems": 1,
               "items": {
                   "type": "object",
                   "properties": {
                       "row": {"type": "integer"},
                       "is_noise": {"type": "string", "enum": ["yes", "no"]},
                       "category": {
                           "type": "string",
                           "enum": ["unavoidable", "avoidable", "unclear", ""],
                           "description": "Empty when is_noise is no"},
                       "quote": {
                           "type": "array", "items": {"type": "string"},
                           "description": "Sentences copied exactly from "
                                          "the review; empty when is_noise "
                                          "is no"},
                       "reason": {"type": "string"},
                   },
                   "required": ["row", "is_noise", "category", "quote",
                                "reason"],
               },
           },
       },
       "required": ["hotel", "items"]})
async def add_judgments_tool(args):
    """工具 add_judgments。"""
    return run(add_judgments, batch_text, args["hotel"], args["items"])


def keyword_text(result):
    """add_keyword 的结果 -> 新命中几条 + 第一批。"""
    return (f"Keyword round {result['round']}: {result['new_hits']} new "
            f"review(s) matched.\n{batch_text(result)}")


@tool("add_keyword",
      "Start the next round with one new keyword (plain text, not case "
      "sensitive). Only allowed when the current round is finished. "
      "Returns the first batch.",
      {"type": "object",
       "properties": {
           "hotel": HOTEL_PARAM,
           "word": {"type": "string"},
           "reasoning": {"type": "string",
                         "description": "Why you chose this keyword"},
       },
       "required": ["hotel", "word", "reasoning"]})
async def add_keyword_tool(args):
    """工具 add_keyword。"""
    return run(add_keyword, keyword_text, args["hotel"], args["word"],
               args["reasoning"])


@tool("build_report",
      "After all 7 keyword rounds are finished: compute the verdict, write "
      "the Excel report and the city summary. Returns 6 lines to reply with.",
      {"type": "object",
       "properties": {
           "hotel": HOTEL_PARAM,
           "inference": {"type": "string",
                         "description": "Your [Inference], in English"},
       },
       "required": ["hotel", "inference"]})
async def build_report_tool(args):
    """工具 build_report。"""
    return run(build_report, report_text, args["hotel"], args["inference"])


TOOLS = [next_batch_tool, add_judgments_tool, add_keyword_tool,
         build_report_tool]
# allowed_tools 和 subagent 的 tools 里要写的全名
TOOL_NAMES = [f"mcp__{SERVER_NAME}__{t.name}" for t in TOOLS]


def create_server():
    """把 4 个工具打包成一个进程内的 MCP 服务器，交给 ClaudeAgentOptions。"""
    return create_sdk_mcp_server(name=SERVER_NAME, tools=TOOLS)
