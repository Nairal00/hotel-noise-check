"""utils/hotel_tools.py 的测试：逻辑层 + @tool 包装层。"""

import asyncio
import csv
import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from hotel_checker_agent.utils import hotel_tools as ht
from hotel_checker_agent.utils.hotel_tools import (
    BATCH,
    CORE_KEYWORD,
    HOTELS,
    LOG_HEADER,
    MAX_ROUND,
    NOT_OK_VERDICT,
    OCC_LIMIT,
    OK_VERDICT,
    REPORT,
    SUMMARY_HEADER,
    TOOL_NAMES,
    UNAV_LIMIT,
    Hotel,
    Review,
    add_judgments,
    add_keyword,
    batch_errors,
    batch_text,
    build_report,
    category_errors,
    create_server,
    current_round,
    get_hotel,
    item_errors,
    keyword_text,
    load_hotels,
    log_rows,
    mark_round,
    next_batch,
    normalize,
    quote_errors,
    quote_found,
    rates,
    report_text,
    run,
    unjudged,
    verdict_of,
    write_summary,
)

NAME = "Tiny Inn"
# 测试酒店的 6 条差评 (行号, 原文)，注释是哪一轮会命中它
COMPLAINTS = [
    (2, "Very noisy street at night"),    # 第 0 轮：nois
    (3, "Thin walls, heard everything"),  # 第 0 轮：wall、hear
    (4, "Breakfast was cold"),            # 一直不命中
    (5, "Loud neighbours"),               # 第 1 轮：loud
    (6, "Music from the bar"),            # 第 2 轮：music
    (7, "The room was small"),            # 第 3 轮：small
]
# 第 1～7 轮的词；后 4 个一条都搜不到
KEYWORDS = ["loud", "music", "small", "snor", "thump", "party", "traffic"]
# 判成噪音的行 -> category；不在这里的行判 no
NOISE = {2: "avoidable", 3: "unavoidable", 5: "unclear", 6: "avoidable"}


def hotel_data(complaints=None, total=10, work_dir="runs/Milan/8.0/Tiny_Inn"):
    """造一份 select_hotels 格式的酒店数据。"""
    if complaints is None:
        complaints = COMPLAINTS
    return {"hotel": NAME, "reviews": total, "work_dir": str(work_dir),
            "complaints": [{"row": r, "text": t} for r, t in complaints]}


def judgment(row, text, category=None):
    """一条合格的判断：category 为 None 时判 no，否则判 yes 并整句引用。"""
    if category is None:
        return {"row": row, "is_noise": "no", "category": "", "quote": [],
                "reason": "not about noise"}
    return {"row": row, "is_noise": "yes", "category": category,
            "quote": [text], "reason": "complains about noise"}


def judge_round(name, noise=None):
    """把本轮剩下的差评全部判断完，返回本轮读完时的 next_batch 结果。"""
    noise = NOISE if noise is None else noise
    while True:
        batch = next_batch(name)
        if batch["done"]:
            return batch
        add_judgments(name, [judgment(v["row"], v["text"], noise.get(v["row"]))
                             for v in batch["reviews"]])


def run_rounds(name, keywords=None):
    """判断完第 0 轮，再依次加词、判断，跑完 keywords 这么多轮。"""
    keywords = KEYWORDS if keywords is None else keywords
    judge_round(name)
    for word in keywords:
        add_keyword(name, word, f"why {word}")
        judge_round(name)


def call_tool(tool, args):
    """直接调用 @tool 包装后的工具函数，返回 (文字, 是否 is_error)。"""
    result = asyncio.run(tool.handler(args))
    return result["content"][0]["text"], result.get("is_error", False)


class HotelTestCase(unittest.TestCase):
    """每个测试前装入测试酒店，测完清空 HOTELS。"""

    def setUp(self):
        load_hotels([hotel_data()])
        self.hotel = HOTELS[NAME]

    def tearDown(self):
        HOTELS.clear()


class TestMarkRound(unittest.TestCase):
    """mark_round：只标还没命中过、而且含关键词的差评。"""

    def test_marks_only_new_matches(self):
        """已经命中过的不改轮次，也不计数。"""
        hotel = Hotel(NAME, Path("x"), 3,
                      [Review(2, "noisy", round=0), Review(3, "noisy"),
                       Review(4, "quiet")], [])
        count = mark_round(hotel, re.compile("noisy"), 1)
        self.assertEqual(count, 1)
        self.assertEqual([v.round for v in hotel.reviews], [0, 1, None])


class TestLoadHotels(HotelTestCase):
    """load_hotels：存进内存，并用核心词标出第 0 轮。"""

    def test_hotel_fields(self):
        """名字、工作目录、评论总数、差评都照抄，第 0 轮是核心词。"""
        self.assertEqual(self.hotel.name, NAME)
        self.assertEqual(self.hotel.work_dir, Path("runs/Milan/8.0/Tiny_Inn"))
        self.assertEqual(self.hotel.total, 10)
        self.assertEqual([(v.row, v.text) for v in self.hotel.reviews],
                         COMPLAINTS)
        self.assertEqual(self.hotel.rounds,
                         [(CORE_KEYWORD, "initial core keywords")])

    def test_round_zero_marked(self):
        """含 nois / wall / hear 的标成第 0 轮，其他的还没命中。"""
        self.assertEqual([v.round for v in self.hotel.reviews],
                         [0, 0, None, None, None, None])

    def test_core_keyword_ignores_case(self):
        """核心词不分大小写，四个词都能命中。"""
        load_hotels([hotel_data([(2, "NOISE"), (3, "WaLL"), (4, "HEAR"),
                                 (5, "Siren"), (6, "fine")])])
        self.assertEqual([v.round for v in HOTELS[NAME].reviews],
                         [0, 0, 0, 0, None])

    def test_old_hotels_cleared(self):
        """再装一次，上一次的酒店不留。"""
        other = dict(hotel_data(), hotel="Other Inn")
        load_hotels([other])
        self.assertEqual(list(HOTELS), ["Other Inn"])


class TestGetHotel(HotelTestCase):
    """get_hotel：名字必须完全一致。"""

    def test_known_name(self):
        """名字对就返回这家酒店。"""
        self.assertIs(get_hotel(NAME), self.hotel)

    def test_unknown_name(self):
        """名字差一个字也报错，错误里带着这个名字。"""
        for name in ("Tiny inn", " Tiny Inn", ""):
            with self.subTest(name=name):
                with self.assertRaisesRegex(ValueError, "Unknown hotel"):
                    get_hotel(name)


class TestNextBatch(HotelTestCase):
    """next_batch：本轮没判断的下一批，或本轮读完的条数。"""

    def test_round_zero_batch(self):
        """第 0 轮有 2 条没判断。"""
        self.assertEqual(next_batch(NAME), {
            "round": 0, "done": False, "left": 2,
            "reviews": [{"row": 2, "text": "Very noisy street at night"},
                        {"row": 3, "text": "Thin walls, heard everything"}]})

    def test_batch_size_limit(self):
        """超过 BATCH 条时一批只给 BATCH 条，left 是全部剩下的条数。"""
        load_hotels([hotel_data([(r, "noisy") for r in range(2, BATCH + 12)])])
        result = next_batch(NAME)
        self.assertEqual(result["left"], BATCH + 10)
        self.assertEqual(len(result["reviews"]), BATCH)
        self.assertEqual(result["reviews"][0]["row"], 2)

    def test_round_with_no_hits(self):
        """本轮一条都没命中：直接算读完，0 条。"""
        load_hotels([hotel_data([(2, "cold breakfast")])])
        self.assertEqual(next_batch(NAME),
                         {"round": 0, "done": True, "read": 0, "noise": 0})

    def test_round_finished(self):
        """本轮判断完：给读了几条、几条是噪音。"""
        add_judgments(NAME, [judgment(2, "Very noisy street at night",
                                      "avoidable"),
                             judgment(3, "Thin walls, heard everything")])
        self.assertEqual(next_batch(NAME),
                         {"round": 0, "done": True, "read": 2, "noise": 1})
        self.assertEqual(unjudged(self.hotel), [])

    def test_unknown_hotel(self):
        """名字不对就报错。"""
        with self.assertRaises(ValueError):
            next_batch("Nope")


class TestQuoteHelpers(unittest.TestCase):
    """normalize 和 quote_found：quote 检查时忽略大小写、标点、多余空格。"""

    def test_normalize(self):
        """转小写，标点变空格，多个空格合成一个。"""
        self.assertEqual(normalize("  Thin,  WALLS!\nheard... "),
                         "thin walls heard")

    def test_quote_found(self):
        """大小写、标点、空格不同也算找到。"""
        text = "The walls are thin; I heard everything!"
        for quote in ("the walls are thin", "WALLS ARE THIN!",
                      "thin I heard"):
            with self.subTest(quote=quote):
                self.assertTrue(quote_found(quote, text))

    def test_quote_not_found(self):
        """原文里没有、空的、只有标点的都不算。"""
        text = "The walls are thin"
        for quote in ("walls were thin", "", "   ", "!!!"):
            with self.subTest(quote=quote):
                self.assertFalse(quote_found(quote, text))


class TestCategoryErrors(unittest.TestCase):
    """category_errors：yes 时三选一，no 时留空。"""

    def test_valid(self):
        """合格的组合没有错误。"""
        for is_noise, category in (("yes", "unavoidable"), ("yes", "avoidable"),
                                   ("yes", "unclear"), ("no", "")):
            with self.subTest(is_noise=is_noise, category=category):
                self.assertEqual(category_errors(is_noise, category), [])

    def test_yes_needs_category(self):
        """yes 时 category 为空或不在三个值里都报错。"""
        for category in ("", "loud", "Unavoidable", None):
            with self.subTest(category=category):
                self.assertEqual(len(category_errors("yes", category)), 1)

    def test_no_needs_empty_category(self):
        """no 时填了 category 报错。"""
        self.assertEqual(category_errors("no", "avoidable"),
                         ["category must be empty when is_noise is 'no'"])

    def test_bad_is_noise_not_checked_here(self):
        """is_noise 不是 yes/no 时这里不管，由 item_errors 报。"""
        self.assertEqual(category_errors("maybe", "loud"), [])


class TestQuoteErrors(unittest.TestCase):
    """quote_errors：quote 是列表；yes 至少一句且出自原文，no 为空。"""

    TEXT = "Thin walls. Street noise all night."

    def test_valid(self):
        """yes 带原文句子、no 为空列表都没问题。"""
        self.assertEqual(quote_errors("yes", ["thin walls", "Street noise"],
                                      self.TEXT), [])
        self.assertEqual(quote_errors("no", [], self.TEXT), [])

    def test_not_a_list_of_strings(self):
        """不是列表、列表里有非文字都报错。"""
        for quote in ("thin walls", None, ["thin walls", 1]):
            with self.subTest(quote=quote):
                self.assertEqual(quote_errors("yes", quote, self.TEXT),
                                 ["quote must be a list of sentences"])

    def test_yes_needs_quote(self):
        """yes 时空列表报错。"""
        self.assertEqual(len(quote_errors("yes", [], self.TEXT)), 1)

    def test_no_needs_empty_quote(self):
        """no 时有 quote 报错。"""
        self.assertEqual(quote_errors("no", ["thin walls"], self.TEXT),
                         ["quote must be empty when is_noise is 'no'"])

    def test_each_missing_quote_reported(self):
        """每句找不到的都单独报一条。"""
        errors = quote_errors("yes", ["thin walls", "loud music", "bad bed"],
                              self.TEXT)
        self.assertEqual(errors,
                         ["quote not found in the review: 'loud music'",
                          "quote not found in the review: 'bad bed'"])


class TestItemErrors(unittest.TestCase):
    """item_errors：一条判断的全部问题，每条带行号。"""

    REVIEW = Review(5, "Thin walls")

    def test_valid_item(self):
        """合格的 yes 和 no 都没有错误。"""
        self.assertEqual(item_errors(judgment(5, "Thin walls", "unavoidable"),
                                     self.REVIEW), [])
        self.assertEqual(item_errors(judgment(5, "Thin walls"), self.REVIEW),
                         [])

    def test_bad_is_noise(self):
        """is_noise 不是 yes/no 报错，错误前面带行号。"""
        item = dict(judgment(5, "Thin walls"), is_noise="maybe")
        self.assertEqual(item_errors(item, self.REVIEW),
                         ["row 5: is_noise must be 'yes' or 'no', "
                          "got 'maybe'"])

    def test_empty_reason(self):
        """reason 缺失、空白、不是文字都报错。"""
        for reason in ("", "   ", None, 3):
            with self.subTest(reason=reason):
                item = dict(judgment(5, "Thin walls"), reason=reason)
                self.assertEqual(item_errors(item, self.REVIEW),
                                 ["row 5: reason is empty"])

    def test_missing_fields(self):
        """只给 row：is_noise、quote、reason 三个错误都报出来。"""
        self.assertEqual(len(item_errors({"row": 5}, self.REVIEW)), 3)

    def test_missing_category_on_no(self):
        """no 时不给 category 键也行，默认当空。"""
        item = judgment(5, "Thin walls")
        del item["category"]
        self.assertEqual(item_errors(item, self.REVIEW), [])


class TestBatchErrors(HotelTestCase):
    """batch_errors：行号必须属于本轮、没判断过、本批不重复。"""

    def test_valid_batch(self):
        """本轮的两条都合格，没有错误。"""
        items = [judgment(2, "Very noisy street at night", "avoidable"),
                 judgment(3, "Thin walls, heard everything", "unavoidable")]
        self.assertEqual(batch_errors(self.hotel, items), [])

    def test_row_sent_twice(self):
        """同一行发两次，第二次报错。"""
        item = judgment(2, "Very noisy street at night", "avoidable")
        self.assertEqual(batch_errors(self.hotel, [item, item]),
                         ["row 2: sent twice in this batch"])

    def test_row_not_in_round(self):
        """不是本轮的行（没命中、根本不存在、没给行号）都报错。"""
        for row in (4, 99, None):
            with self.subTest(row=row):
                self.assertEqual(batch_errors(self.hotel, [judgment(row, "x")]),
                                 [f"row {row}: not a review of round 0"])

    def test_already_judged(self):
        """已经判断过的行再发报错。"""
        item = judgment(2, "Very noisy street at night", "avoidable")
        add_judgments(NAME, [item])
        self.assertEqual(batch_errors(self.hotel, [item]),
                         ["row 2: already judged"])

    def test_all_errors_collected(self):
        """一批里的错误全部列出，不是遇到第一个就停。"""
        items = [dict(judgment(2, "x"), is_noise="maybe"),
                 judgment(3, "Thin walls", "loud"),
                 judgment(99, "x")]
        self.assertEqual(len(batch_errors(self.hotel, items)), 3)


class TestAddJudgments(HotelTestCase):
    """add_judgments：整批检查通过才存，存完返回下一批。"""

    def test_saves_and_returns_next(self):
        """存下判断，reason 去掉首尾空格，返回本轮读完的结果。"""
        result = add_judgments(NAME, [
            dict(judgment(2, "Very noisy street at night", "avoidable"),
                 reason="  street  "),
            judgment(3, "Thin walls, heard everything")])
        self.assertEqual(result,
                         {"round": 0, "done": True, "read": 2, "noise": 1})
        noisy, quiet = self.hotel.reviews[:2]
        self.assertEqual((noisy.is_noise, noisy.category, noisy.quote,
                          noisy.reason),
                         ("yes", "avoidable", ["Very noisy street at night"],
                          "street"))
        self.assertEqual((quiet.is_noise, quiet.category, quiet.quote),
                         ("no", "", []))

    def test_partial_batch_returns_rest(self):
        """只判断一条时，返回剩下的那条。"""
        result = add_judgments(NAME, [judgment(2, "x")])
        self.assertEqual(result["left"], 1)
        self.assertEqual(result["reviews"][0]["row"], 3)

    def test_quote_is_copied(self):
        """存的是 quote 的副本，之后改原列表不影响已存的。"""
        item = judgment(2, "Very noisy street at night", "avoidable")
        add_judgments(NAME, [item])
        item["quote"].append("changed")
        self.assertEqual(self.hotel.reviews[0].quote,
                         ["Very noisy street at night"])

    def test_one_bad_item_saves_nothing(self):
        """有一条不对，整批都不存，错误里列出问题。"""
        items = [judgment(2, "Very noisy street at night", "avoidable"),
                 judgment(3, "Thin walls", "loud")]
        with self.assertRaisesRegex(ValueError, "Nothing was saved"):
            add_judgments(NAME, items)
        self.assertEqual(len(unjudged(self.hotel)), 2)

    def test_empty_items(self):
        """空列表报错。"""
        with self.assertRaisesRegex(ValueError, "items is empty"):
            add_judgments(NAME, [])


class TestAddKeyword(HotelTestCase):
    """add_keyword：上一轮读完才能加词，最多 7 轮。"""

    def test_round_not_finished(self):
        """第 0 轮还没判断完，不许加词。"""
        with self.assertRaisesRegex(ValueError, "Round 0 is not finished"):
            add_keyword(NAME, "loud", "why")
        self.assertEqual(current_round(self.hotel), 0)

    def test_adds_round(self):
        """加词后进入第 1 轮，只命中还没命中过的差评，返回第一批。"""
        judge_round(NAME)
        result = add_keyword(NAME, "  LOUD ", "  people  ")
        self.assertEqual(current_round(self.hotel), 1)
        self.assertEqual(self.hotel.rounds[1], ("LOUD", "people"))
        self.assertEqual(result, {
            "new_hits": 1, "round": 1, "done": False, "left": 1,
            "reviews": [{"row": 5, "text": "Loud neighbours"}]})

    def test_old_hits_not_counted(self):
        """第 0 轮命中过的不再算新命中。"""
        judge_round(NAME)
        result = add_keyword(NAME, "noisy", "why")
        self.assertEqual(result["new_hits"], 0)
        self.assertTrue(result["done"])
        self.assertEqual(self.hotel.reviews[0].round, 0)

    def test_special_chars_are_plain_text(self):
        """词里的 . 等符号按普通文字匹配，不当正则。"""
        load_hotels([hotel_data([(2, "cold"), (3, "a.c broken"),
                                 (4, "abc broken")])])
        judge_round(NAME)
        add_keyword(NAME, "a.c", "air conditioning")
        self.assertEqual([v.round for v in HOTELS[NAME].reviews],
                         [None, 1, None])

    def test_empty_word_or_reasoning(self):
        """词或理由为空（含只有空格）都报错，也不进入新一轮。"""
        judge_round(NAME)
        for word, reasoning, message in (("", "why", "word is empty"),
                                         ("  ", "why", "word is empty"),
                                         ("loud", " ", "reasoning is empty")):
            with self.subTest(word=word, reasoning=reasoning):
                with self.assertRaisesRegex(ValueError, message):
                    add_keyword(NAME, word, reasoning)
        self.assertEqual(current_round(self.hotel), 0)

    def test_max_rounds(self):
        """跑满 7 轮后不许再加词。"""
        run_rounds(NAME)
        self.assertEqual(current_round(self.hotel), MAX_ROUND)
        with self.assertRaisesRegex(ValueError, "All 7 keyword rounds"):
            add_keyword(NAME, "extra", "why")


class TestRatesAndVerdict(unittest.TestCase):
    """rates 和 verdict_of：两个比例和结论。"""

    @staticmethod
    def reviews(categories):
        """按 category 造已判断的评论；None 表示判 no。"""
        return [Review(i, "x", round=0,
                       is_noise="no" if c is None else "yes",
                       category=c or "")
                for i, c in enumerate(categories)]

    def test_rates(self):
        """噪音条数 / 总数；unavoidable 条数 / 噪音条数。"""
        reviews = self.reviews(["unavoidable", "avoidable", "unclear",
                                "avoidable", None])
        self.assertEqual(rates(reviews, 10), (0.4, 0.25))

    def test_no_noise(self):
        """没有噪音差评时两个比例都是 0。"""
        self.assertEqual(rates(self.reviews([None, None]), 10), (0.0, 0.0))
        self.assertEqual(rates([], 10), (0.0, 0.0))

    def test_unjudged_not_counted(self):
        """还没判断的评论不算噪音。"""
        self.assertEqual(rates([Review(2, "noisy", round=0)], 10), (0.0, 0.0))

    def test_verdict(self):
        """两个比例都严格低于门槛才能订；等于门槛也不行。"""
        cases = [
            (0.0, 0.0, OK_VERDICT),
            (OCC_LIMIT - 0.001, UNAV_LIMIT - 0.001, OK_VERDICT),
            (OCC_LIMIT, 0.0, NOT_OK_VERDICT),
            (0.0, UNAV_LIMIT, NOT_OK_VERDICT),
            (0.5, 0.9, NOT_OK_VERDICT),
        ]
        for occ, unav, expected in cases:
            with self.subTest(occ=occ, unav=unav):
                self.assertEqual(verdict_of(occ, unav), expected)


class TestLogRows(HotelTestCase):
    """log_rows：每轮一行，比例算到这一轮为止。"""

    def test_rows_after_all_rounds(self):
        """8 行（第 0～7 轮），比例逐轮累计。"""
        run_rounds(NAME)
        rows = log_rows(self.hotel)
        self.assertEqual(len(rows), MAX_ROUND + 1)
        self.assertEqual(rows[0], [0, CORE_KEYWORD, "initial core keywords",
                                   2, 2, 0.2, 0.5, "2, 3", "2, 3"])
        self.assertEqual(rows[1][:7], [1, "loud", "why loud", 1, 1, 0.3, 1 / 3])
        self.assertEqual(rows[2][5:], [0.4, 0.25, "6", "6"])
        # 第 3 轮命中了一条但不是噪音
        self.assertEqual(rows[3][3:], [1, 0, 0.4, 0.25, "7", ""])
        for row in rows[4:]:
            self.assertEqual(row[3:], [0, 0, 0.4, 0.25, "", ""])


class TestBuildReport(HotelTestCase):
    """build_report：跑满 7 轮、判断完、写了 inference 才能出报告。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()  # pylint: disable=consider-using-with
        self.city_dir = Path(self.tmp.name) / "Milan"
        self.work_dir = self.city_dir / "8.0" / "Tiny_Inn"
        load_hotels([hotel_data(work_dir=self.work_dir)])
        self.hotel = HOTELS[NAME]

    def tearDown(self):
        super().tearDown()
        self.tmp.cleanup()

    def test_needs_all_rounds(self):
        """没跑满 7 轮报错，也不写文件。"""
        run_rounds(NAME, KEYWORDS[:3])
        with self.assertRaisesRegex(ValueError, "Only 3 of 7"):
            build_report(NAME, "inference")
        self.assertFalse(self.city_dir.exists())

    def test_last_round_unfinished(self):
        """第 7 轮还有没判断的，报错。"""
        run_rounds(NAME, KEYWORDS[:6])
        add_keyword(NAME, "breakfast", "why")
        with self.assertRaisesRegex(ValueError, "Round 7 is not finished"):
            build_report(NAME, "inference")

    def test_empty_inference(self):
        """inference 为空或只有空格，报错。"""
        run_rounds(NAME)
        with self.assertRaisesRegex(ValueError, "inference is empty"):
            build_report(NAME, "   ")

    def test_result(self):
        """返回 6 项：4/10 是噪音，1/4 是 unavoidable，不能订。"""
        run_rounds(NAME)
        result = build_report(NAME, " Street noise. ")
        self.assertEqual(result, {
            "hotel": NAME, "verdict": NOT_OK_VERDICT,
            "occurrence_rate": 0.4, "unavoidable_rate": 0.25,
            "rounds": MAX_ROUND + 1, "report": str(self.work_dir / REPORT)})

    def test_ok_verdict(self):
        """总评论数多时比例低，可以订。"""
        load_hotels([hotel_data(total=100, work_dir=self.work_dir)])
        run_rounds(NAME)
        result = build_report(NAME, "Little noise.")
        self.assertEqual(result["verdict"], OK_VERDICT)
        self.assertEqual(result["occurrence_rate"], 0.04)

    def test_xlsx_sheets(self):
        """Excel 有 4 个工作表，summary 表的数字和结果一致。"""
        run_rounds(NAME)
        build_report(NAME, " Street noise. ")
        wb = load_workbook(self.work_dir / REPORT)
        self.assertEqual(wb.sheetnames,
                         ["summary", "log", "judgments",
                          "all negative reviews"])
        summary = {r[0].value: r[1] for r in wb["summary"].iter_rows(min_row=2)}
        self.assertEqual(
            {k: c.value for k, c in summary.items()},
            {"hotel": NAME, "verdict": NOT_OK_VERDICT, "total_reviews": 10,
             "noise_reviews": 4, "unavoidable_reviews": 1,
             "occurrence_rate": 0.4, "unavoidable_rate": 0.25,
             "keywords_used": "; ".join([CORE_KEYWORD, *KEYWORDS]),
             "inference": "Street noise."})
        self.assertEqual(summary["occurrence_rate"].number_format, "0.0%")
        self.assertEqual(summary["unavoidable_rate"].number_format, "0.0%")
        self.assertEqual(summary["hotel"].number_format, "General")

    def test_xlsx_log_sheet(self):
        """log 表：表头加粗、冻结，8 行数据，比例列是百分比格式。"""
        run_rounds(NAME)
        build_report(NAME, "x")
        ws = load_workbook(self.work_dir / REPORT)["log"]
        rows = list(ws.values)
        self.assertEqual(list(rows[0]), LOG_HEADER)
        self.assertEqual(len(rows), MAX_ROUND + 2)
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws["F2"].number_format, "0.0%")
        self.assertEqual(ws.column_dimensions["C"].width, 60)
        self.assertTrue(ws["C2"].alignment.wrap_text)
        self.assertEqual(ws.column_dimensions["A"].width, 18)

    def test_xlsx_review_sheets(self):
        """judgments 只有判断过的，按行号排；all negative reviews 是全部差评。"""
        run_rounds(NAME)
        build_report(NAME, "x")
        wb = load_workbook(self.work_dir / REPORT)
        judged = list(wb["judgments"].values)[1:]
        self.assertEqual([r[0] for r in judged], [2, 3, 5, 6, 7])
        self.assertEqual(judged[0], (2, 0, "Very noisy street at night", "yes",
                                     "avoidable", "Very noisy street at night",
                                     "complains about noise"))
        self.assertEqual(judged[-1][3:6], ("no", None, None))
        everything = list(wb["all negative reviews"].values)[1:]
        self.assertEqual([(r[0], r[1]) for r in everything],
                         [(2, "Y"), (3, "Y"), (4, None), (5, None),
                          (6, None), (7, None)])

    def test_summary_csv(self):
        """summary.csv 在城市文件夹，report 写相对路径。"""
        run_rounds(NAME)
        build_report(NAME, "x")
        with open(self.city_dir / "summary.csv", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        self.assertEqual(rows, [
            SUMMARY_HEADER,
            ["Milan", "8.0", NAME, NOT_OK_VERDICT, "0.4", "0.25", "8",
             "8.0/Tiny_Inn/" + REPORT]])

    def test_report_rebuilt(self):
        """同一家再出一次报告，覆盖 xlsx，summary.csv 里还是只有一行。"""
        run_rounds(NAME)
        build_report(NAME, "first")
        build_report(NAME, "second")
        with open(self.city_dir / "summary.csv", encoding="utf-8") as f:
            self.assertEqual(len(list(csv.reader(f))), 2)
        summary = load_workbook(self.work_dir / REPORT)["summary"]
        self.assertEqual(summary["B10"].value, "second")


class TestWriteSummary(unittest.TestCase):
    """write_summary：每家酒店一行，同一家的旧行替换掉。"""

    RESULT = {"verdict": OK_VERDICT, "occurrence_rate": 0.123456,
              "unavoidable_rate": 1 / 3, "rounds": 8}

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()  # pylint: disable=consider-using-with
        self.city_dir = Path(self.tmp.name) / "Milan"
        self.path = self.city_dir / "summary.csv"

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, name):
        """给 name 这家酒店写一行，返回 summary.csv 的全部行。"""
        work_dir = self.city_dir / "8.0" / name.replace(" ", "_")
        work_dir.mkdir(parents=True, exist_ok=True)
        hotel = Hotel(name, work_dir, 10, [], [])
        write_summary(hotel, self.RESULT, work_dir / REPORT)
        with open(self.path, encoding="utf-8") as f:
            return list(csv.reader(f))

    def test_rates_rounded(self):
        """两个比例保留 4 位小数。"""
        rows = self.write("A Inn")
        self.assertEqual(rows[1][4:6], ["0.1235", "0.3333"])

    def test_other_hotels_kept(self):
        """其他酒店的行留着，同一家的旧行换成新的放最后。"""
        self.write("A Inn")
        self.write("B Inn")
        rows = self.write("A Inn")
        self.assertEqual([r[2] for r in rows], ["hotel", "B Inn", "A Inn"])

    def test_blank_lines_dropped(self):
        """已有文件里的空行去掉。"""
        self.city_dir.mkdir(parents=True)
        self.path.write_text(",".join(SUMMARY_HEADER) + "\n\n", encoding="utf-8")
        rows = self.write("A Inn")
        self.assertEqual(len(rows), 2)


class TestToolText(unittest.TestCase):
    """把逻辑层结果转成 AI 读的文字。"""

    def test_batch_text_reviews(self):
        """还有没判断的：列出这一批，提示用 add_judgments。"""
        text = batch_text({"round": 0, "done": False, "left": 3,
                           "reviews": [{"row": 2, "text": "noisy"}]})
        self.assertEqual(text.splitlines(), [
            "Round 0: 3 review(s) left to judge. This batch has 1:",
            "row 2: noisy",
            "Judge every review above, then send them all with add_judgments."])

    def test_batch_text_round_done(self):
        """没满 7 轮：提示加下一个词。"""
        text = batch_text({"round": 2, "done": True, "read": 4, "noise": 1})
        self.assertEqual(text, "Round 2 finished: read 4 review(s), 1 judged "
                               "as noise. Next: call add_keyword (round 3 of 7).")

    def test_batch_text_all_done(self):
        """满 7 轮：提示出报告。"""
        text = batch_text({"round": MAX_ROUND, "done": True, "read": 0,
                           "noise": 0})
        self.assertTrue(text.endswith("Next: write the inference and call "
                                      "build_report."))

    def test_keyword_text(self):
        """先说新命中几条，再接 batch_text。"""
        text = keyword_text({"new_hits": 0, "round": 1, "done": True,
                             "read": 0, "noise": 0})
        self.assertTrue(text.startswith(
            "Keyword round 1: 0 new review(s) matched.\nRound 1 finished"))

    def test_report_text(self):
        """6 行，比例显示成百分比。"""
        text = report_text({"hotel": NAME, "verdict": OK_VERDICT,
                            "occurrence_rate": 0.04, "unavoidable_rate": 0.25,
                            "rounds": 8, "report": "/r.xlsx"})
        self.assertEqual(text.splitlines(), [
            f"hotel: {NAME}", f"verdict: {OK_VERDICT}",
            "occurrence_rate: 4.0%", "unavoidable_rate: 25.0%",
            "rounds: 8", "report: /r.xlsx"])


class TestRun(unittest.TestCase):
    """run：ValueError 变成 is_error，其他异常照常抛出。"""

    def test_success(self):
        """成功时只有 content，没有 is_error。"""
        self.assertEqual(run(lambda x: x * 2, str, 21),
                         {"content": [{"type": "text", "text": "42"}]})

    def test_value_error(self):
        """ValueError 的内容原样给 AI，标 is_error。"""
        def fail():
            raise ValueError("bad input")
        self.assertEqual(run(fail, str),
                         {"content": [{"type": "text", "text": "bad input"}],
                          "is_error": True})

    def test_other_errors_raised(self):
        """不是 ValueError 的异常不吞掉。"""
        def fail():
            raise KeyError("x")
        with self.assertRaises(KeyError):
            run(fail, str)


class TestTools(HotelTestCase):
    """4 个 @tool 工具：调用逻辑层，返回文字。"""

    def test_tool_names(self):
        """工具全名是 mcp__hotel__<工具名>。"""
        self.assertEqual(TOOL_NAMES, ["mcp__hotel__next_batch",
                                      "mcp__hotel__add_judgments",
                                      "mcp__hotel__add_keyword",
                                      "mcp__hotel__build_report"])

    def test_create_server(self):
        """服务器名是 hotel。"""
        server = create_server()
        self.assertEqual((server["type"], server["name"]), ("sdk", "hotel"))

    def test_next_batch_tool(self):
        """next_batch 返回第一批的文字。"""
        text, is_error = call_tool(ht.next_batch_tool, {"hotel": NAME})
        self.assertFalse(is_error)
        self.assertIn("row 2: Very noisy street at night", text)

    def test_unknown_hotel_is_error(self):
        """名字不对：is_error，文字是错误原因。"""
        text, is_error = call_tool(ht.next_batch_tool, {"hotel": "Nope"})
        self.assertTrue(is_error)
        self.assertIn("Unknown hotel", text)

    def test_add_judgments_tool(self):
        """add_judgments 存下判断，返回本轮读完。"""
        items = [judgment(2, "Very noisy street at night", "avoidable"),
                 judgment(3, "Thin walls, heard everything")]
        text, is_error = call_tool(ht.add_judgments_tool,
                                   {"hotel": NAME, "items": items})
        self.assertFalse(is_error)
        self.assertTrue(text.startswith("Round 0 finished"))

    def test_add_judgments_tool_error(self):
        """一批不合格：is_error，列出问题。"""
        text, is_error = call_tool(ht.add_judgments_tool,
                                   {"hotel": NAME, "items": [judgment(9, "x")]})
        self.assertTrue(is_error)
        self.assertIn("row 9: not a review of round 0", text)

    def test_add_keyword_tool(self):
        """add_keyword 返回新命中条数和第一批。"""
        judge_round(NAME)
        text, is_error = call_tool(ht.add_keyword_tool,
                                   {"hotel": NAME, "word": "loud",
                                    "reasoning": "people"})
        self.assertFalse(is_error)
        self.assertIn("Keyword round 1: 1 new review(s) matched.", text)
        self.assertIn("row 5: Loud neighbours", text)

    def test_build_report_tool(self):
        """build_report 返回 6 行。"""
        tmp = tempfile.TemporaryDirectory()  # pylint: disable=consider-using-with
        self.addCleanup(tmp.cleanup)
        load_hotels([hotel_data(work_dir=Path(tmp.name) / "Milan" / "8.0"
                                / "Tiny_Inn")])
        run_rounds(NAME)
        text, is_error = call_tool(ht.build_report_tool,
                                   {"hotel": NAME, "inference": "Street noise."})
        self.assertFalse(is_error)
        self.assertEqual(len(text.splitlines()), 6)
        self.assertIn("occurrence_rate: 40.0%", text)


if __name__ == "__main__":
    unittest.main()
