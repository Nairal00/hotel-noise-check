"""第 5 步：下结论，生成 hotel_noise_analysis_log.xlsx。

用法: python scripts/build_report.py --hotel "酒店名"
"""
import argparse
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

REPORT = 'hotel_noise_analysis_log.xlsx'
# 第 5 步 agent 写的 [Inference]
INFERENCE = 'inference.txt'
# 判定规则（用户定的）
OCC_LIMIT, UNAV_LIMIT = 0.2, 0.3


def read(path):
    """把一个 CSV 读成一行行的字典。"""
    with open(path, encoding='utf-8') as f:
        return list(csv.DictReader(f))


def sheet(ws, header, data, pct=(), wide=()):
    """写一个工作表：表头加粗，比例显示成百分比，长文字列加宽并换行。"""
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for d in data:
        ws.append(d)
    for i, h in enumerate(header, 1):
        # 列字母，如 A、B
        col = ws.cell(1, i).column_letter
        ws.column_dimensions[col].width = 60 if h in wide else 18
        for c in ws[col][1:]:
            if h in pct:
                c.number_format = '0.0%'
            if h in wide:
                c.alignment = Alignment(wrap_text=True, vertical='top')
    # 往下滚动时表头不动
    ws.freeze_panes = 'A2'


def main():
    """读中间文件，生成四个工作表的 Excel。"""
    parser = argparse.ArgumentParser()
    parser.add_argument('--hotel', required=True, help='exact Hotel_Name')
    args = parser.parse_args()
    # 第 1 步生成的文件，缺了就提示先跑 prepare.py
    for path in ('total.txt', 'judgments.csv', 'todo.csv'):
        if not os.path.exists(path):
            sys.exit(f'{path} not found: run scripts/prepare.py first')
    # log.csv 是每轮的记录，没有或只有表头都说明还没记过一轮
    log = read('log.csv') if os.path.exists('log.csv') else []
    if not log:
        sys.exit('log.csv has no rounds: run scripts/log_round.py first')

    # 没有 [Inference] 就不生成报告
    if not os.path.exists(INFERENCE):
        sys.exit(f'{INFERENCE} not found: write the [Inference] first')
    with open(INFERENCE, encoding='utf-8') as f:
        inference = f.read().strip()
    with open('total.txt', encoding='utf-8') as f:
        total = int(f.read())
    judg, todo = read('judgments.csv'), read('todo.csv')
    # 行号 → 差评原文，给 judgments 页配原文
    text = {t['row']: t['text'] for t in todo}
    # judgments 里的每个行号都要能在 todo.csv 找到原文
    unknown = sorted({j['row'] for j in judg} - text.keys(), key=int)
    if unknown:
        sys.exit(f'rows {unknown} in judgments.csv are not in todo.csv')
    # 最后一轮 = 最终指标
    last = log[-1]
    occ = float(last['occurrence_rate'])
    unav = float(last['unavoidable_rate'])
    if occ < OCC_LIMIT and unav < UNAV_LIMIT:
        verdict = 'OK to book'
    else:
        verdict = 'Not suitable for noise-sensitive guests'

    wb = Workbook()
    # 第 1 页：总结
    ws = wb.active
    ws.title = 'summary'
    sheet(ws, ['item', 'value'], [
        ['hotel', args.hotel],
        ['verdict', verdict],
        ['total_reviews', total],
        ['noise_reviews', sum(j['is_noise'] == 'yes' for j in judg)],
        ['occurrence_rate', occ],
        ['unavoidable_rate', unav],
        ['keywords_used', '; '.join(l['keyword'] for l in log)],
        ['inference', inference],
    ], wide=('value',))
    # 两个指标显示成百分比
    ws['B6'].number_format = ws['B7'].number_format = '0.0%'

    # 第 2 页：每轮记录
    sheet(wb.create_sheet('log'), list(log[0].keys()),
          [[int(l['round']), l['keyword'], l['reasoning'],
            int(l['new_hits_current_round']),
            int(l['new_positive_noise_reviews']),
            float(l['occurrence_rate']), float(l['unavoidable_rate']),
            l['hit_rows'], l['noise_rows']] for l in log],
          pct=('occurrence_rate', 'unavoidable_rate'),
          wide=('reasoning', 'hit_rows', 'noise_rows'))

    # 第 3 页：每条已读差评的判断，附原文，按原数据行号排序
    sheet(wb.create_sheet('judgments'),
          ['row', 'round', 'text', 'is_noise', 'category', 'quote', 'reason'],
          [[int(j['row']), int(j['round']), text[j['row']], j['is_noise'],
            j['category'], j['quote'], j['reason']]
           for j in sorted(judg, key=lambda j: int(j['row']))],
          wide=('text', 'quote', 'reason'))

    # 第 4 页：所有写了差评的评论
    sheet(wb.create_sheet('all negative reviews'),
          ['row', 'query_hit', 'text'],
          [[int(t['row']), t['query_hit'], t['text']] for t in todo],
          wide=('text',))

    wb.save(REPORT)
    # 报告结论和两个指标
    print(verdict, f'{occ:.1%}', f'{unav:.1%}')


if __name__ == '__main__':
    main()
